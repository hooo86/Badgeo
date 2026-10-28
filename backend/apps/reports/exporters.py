import csv
import io
from datetime import datetime

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BLUE = "004AC6"


def _csv_safe(v):
    """Neutralise spreadsheet formula injection in exported text."""
    return "'" + v if isinstance(v, str) and v[:1] in "=+-@" and len(v) > 1 and not v[1:].replace(".", "").isdigit() else v


def to_csv(title, headers, rows, meta):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(headers)
    for r in rows:
        w.writerow([_csv_safe(c) for c in r])
    return buf.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", "csv"


def to_xlsx(title, headers, rows, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14, color=BLUE)
    ws.append([meta])
    ws.append([])
    ws.append(headers)
    for c in ws[4]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([_csv_safe(c) for c in r])
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
    ws.freeze_panes = "A5"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"


def to_pdf(title, headers, rows, meta):
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28,
                            title=title)
    st = getSampleStyleSheet()
    cell = st["BodyText"].clone("cell", fontSize=7, leading=9)
    data = [headers] + [[Paragraph(str(c), cell) for c in r] for r in rows]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BLUE)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, 0), 7.5), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f4f6")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c3c6d7")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))

    def footer(canvas, d):
        canvas.setFont("Helvetica", 7)
        canvas.drawString(24, 16, f"Badgeo – {title} – généré le {datetime.now():%d/%m/%Y %H:%M}")
        canvas.drawRightString(landscape(A4)[0] - 24, 16, f"Page {d.page}")

    doc.build([Paragraph(f"<font color='#{BLUE}'><b>{title}</b></font>", st["Title"]), Paragraph(meta, st["Normal"]),
               Spacer(1, 8), t], onFirstPage=footer, onLaterPages=footer)
    return out.getvalue(), "application/pdf", "pdf"


EXPORTERS = {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf}


def respond(fmt, name, title, headers, rows, meta):
    body, ctype, ext = EXPORTERS[fmt](title, headers, rows, meta)
    r = HttpResponse(body, content_type=ctype)
    r["Content-Disposition"] = f'attachment; filename="{name}-{timezone.localdate():%Y%m%d}.{ext}"'
    return r
