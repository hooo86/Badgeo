"""Employee import/export (CSV or XLSX). Columns (French headers, accents optional):
matricule, nom, prenom, unite (ex: "Direction RH > Service Paie"), fonction, date_recrutement,
statut (actif|suspendu|sorti), pointeuse, id_pointeuse."""
import csv
import io
import unicodedata
from datetime import date, datetime

from django.db import transaction
from openpyxl import Workbook, load_workbook

from apps.attendance.services import link_punches
from apps.devices.models import Device
from apps.orgs.models import OrgUnit

from .models import DeviceUserMap, Employee

HEADERS = ["matricule", "nom", "prenom", "unite", "fonction", "date_recrutement", "statut", "pointeuse", "id_pointeuse"]
STATUS_IN = {"actif": "active", "active": "active", "suspendu": "suspended", "suspended": "suspended",
             "sorti": "left", "left": "left", "": "active"}
STATUS_OUT = {"active": "actif", "suspended": "suspendu", "left": "sorti"}
KINDS = ["direction", "service", "department"]


def _norm(h):
    return unicodedata.normalize("NFKD", str(h or "")).encode("ascii", "ignore").decode().strip().lower().replace(" ", "_")


def read_rows(upload):
    raw = upload.read()
    if upload.name.lower().endswith((".xlsx", ".xlsm")):
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        text = raw.decode("utf-8-sig")
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t") if text else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return []
    head = [_norm(h) for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f"date invalide: {v!r}")


def _unit(path):
    parent = None
    for depth, name in enumerate(p.strip() for p in str(path).split(">") if p.strip()):
        parent, _ = OrgUnit.objects.get_or_create(parent=parent, name=name,
                                                  defaults={"kind": KINDS[min(depth, len(KINDS) - 1)]})
    return parent


def import_employees(upload, dry_run=False):
    created = updated = 0
    errors = []
    for i, row in enumerate(read_rows(upload), start=2):
        try:
            with transaction.atomic():
                mat = str(row.get("matricule") or "").strip()
                if not mat or not row.get("nom"):
                    raise ValueError("matricule et nom obligatoires")
                status = STATUS_IN.get(_norm(row.get("statut")))
                if status is None:
                    raise ValueError(f"statut inconnu: {row.get('statut')!r}")
                emp, is_new = Employee.objects.get_or_create(matricule=mat, defaults={
                    "last_name": str(row["nom"]).strip(), "first_name": str(row.get("prenom") or "").strip()})
                emp.last_name, emp.first_name = str(row["nom"]).strip(), str(row.get("prenom") or "").strip()
                emp.job_title, emp.status = str(row.get("fonction") or "").strip(), status
                emp.hire_date = _date(row.get("date_recrutement"))
                if row.get("unite"):
                    emp.org_unit = _unit(row["unite"])
                emp.save()
                if row.get("pointeuse") and row.get("id_pointeuse") not in (None, ""):
                    dev = Device.objects.get(name=str(row["pointeuse"]).strip())
                    uid = str(row["id_pointeuse"]).strip()
                    DeviceUserMap.objects.update_or_create(device=dev, device_user_id=uid, defaults={"employee": emp})
                    link_punches(emp, dev, uid)
                if dry_run:
                    transaction.set_rollback(True)
                created, updated = created + is_new, updated + (not is_new)
        except Exception as e:  # row-level failure must not stop the whole file
            errors.append({"row": i, "error": str(e)})
    return {"created": created, "updated": updated, "errors": errors, "dry_run": dry_run}


def export_rows():
    maps = {}
    for m in DeviceUserMap.objects.select_related("device"):
        maps.setdefault(m.employee_id, m)
    def path(u):
        out = []
        while u:
            out.append(u.name)
            u = u.parent
        return " > ".join(reversed(out))
    for e in Employee.objects.select_related("org_unit"):
        m = maps.get(e.pk)
        yield [e.matricule, e.last_name, e.first_name, path(e.org_unit), e.job_title,
               e.hire_date.isoformat() if e.hire_date else "", STATUS_OUT[e.status],
               m.device.name if m else "", m.device_user_id if m else ""]


def export_csv():
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(HEADERS)
    w.writerows(export_rows())
    return buf.getvalue().encode("utf-8-sig")


def export_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in export_rows():
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
