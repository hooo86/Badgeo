import io
from datetime import date, datetime, time, timedelta

from django.contrib.auth.models import User
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.authtoken.models import Token

from apps.accounts.models import Profile
from apps.attendance.models import PunchEvent
from apps.attendance.services import recompute_day
from apps.corrections.models import Correction
from apps.devices.models import Device
from apps.employees.models import DeviceUserMap, Employee
from apps.orgs.models import OrgUnit
from apps.schedules.models import Shift, ShiftAssignment

J = "application/json"


def auth(username, role, scopes=(), employee=None):
    u = User.objects.create_user(username, password="Str0ng-pass-123")
    p = Profile.objects.create(user=u, role=role)
    p.scopes.set(scopes)
    if employee:
        employee.user = u
        employee.save()
    return {"HTTP_AUTHORIZATION": f"Token {Token.objects.create(user=u).key}"}


class Base(TestCase):
    def setUp(self):
        cache.clear()
        self.unit = OrgUnit.objects.create(kind="direction", name="Prod")
        self.other = OrgUnit.objects.create(kind="direction", name="Compta")
        self.emp = Employee.objects.create(matricule="E1", first_name="Sophie", last_name="Garnier", org_unit=self.unit)
        self.emp2 = Employee.objects.create(matricule="E2", first_name="Karim", last_name="B", org_unit=self.other)
        self.shift = Shift.objects.create(name="Jour", start_time=time(8), end_time=time(17), work_days=[0, 1, 2, 3, 4])
        for e in (self.emp, self.emp2):
            ShiftAssignment.objects.create(employee=e, shift=self.shift, valid_from=date(2026, 1, 1))
        self.d = date(2026, 3, 2)

    def punch(self, emp, h, m=0, d=None):
        d = d or self.d
        return PunchEvent.objects.create(employee=emp, source="manual", timestamp=timezone.make_aware(datetime.combine(d, time(h, m))))


class CorrectionTests(Base):
    def test_workflow_four_eyes_reason_and_raw_preserved(self):
        self.punch(self.emp, 8)  # missing exit
        day = recompute_day(self.emp, self.d)
        self.assertEqual(day.status, "incomplete")
        author, reviewer = auth("hrs", "hr_staff"), auth("hrm", "hr_manager")
        ts = timezone.make_aware(datetime.combine(self.d, time(17))).isoformat()
        body = {"employee": self.emp.pk, "work_date": str(self.d), "kind": "add_punch", "new_timestamp": ts}
        self.assertEqual(self.client.post("/api/corrections/", {**body, "reason": ""}, content_type=J, **author).status_code, 400)
        r = self.client.post("/api/corrections/", {**body, "reason": "Oubli de badgeage confirmé par le chef"}, content_type=J, **author)
        self.assertEqual(r.status_code, 201)
        cid = r.json()["id"]
        self.assertEqual(r.json()["initial"]["status"], "incomplete")
        # author cannot validate; a staff member cannot validate at all
        self.assertEqual(self.client.post(f"/api/corrections/{cid}/validate/", {}, content_type=J, **author).status_code, 403)
        raw_before = PunchEvent.objects.count()
        ok = self.client.post(f"/api/corrections/{cid}/validate/", {"note": "OK"}, content_type=J, **reviewer)
        self.assertEqual(ok.status_code, 200)
        c = Correction.objects.get(pk=cid)
        self.assertEqual((c.state, c.after["status"], c.initial["status"]), ("validated", "present", "incomplete"))
        self.assertEqual(PunchEvent.objects.count(), raw_before + 1)           # a new manual punch, nothing overwritten
        self.assertEqual(self.client.post(f"/api/corrections/{cid}/validate/", {}, content_type=J, **reviewer).status_code, 400)
        self.assertEqual(self.client.patch(f"/api/corrections/{cid}/", {"reason": "x"}, content_type=J, **reviewer).status_code, 405)

    def test_same_user_cannot_validate_own_correction(self):
        u = auth("hrm", "hr_manager")
        r = self.client.post("/api/corrections/", {"employee": self.emp.pk, "work_date": str(self.d), "kind": "set_status",
                             "new_status": "sick", "reason": "Certificat reçu"}, content_type=J, **u)
        self.assertEqual(self.client.post(f"/api/corrections/{r.json()['id']}/validate/", {}, content_type=J, **u).status_code, 403)

    def test_status_override_survives_recompute(self):
        recompute_day(self.emp, self.d)
        hr, mgr = auth("hrs", "hr_staff"), auth("hrm", "hr_manager")
        r = self.client.post("/api/corrections/", {"employee": self.emp.pk, "work_date": str(self.d), "kind": "set_status",
                             "new_status": "mission", "reason": "Mission chez le client"}, content_type=J, **hr)
        self.client.post(f"/api/corrections/{r.json()['id']}/validate/", {}, content_type=J, **mgr)
        self.assertEqual(recompute_day(self.emp, self.d).status, "mission")

    def test_reject_requires_note(self):
        hr, mgr = auth("hrs", "hr_staff"), auth("hrm", "hr_manager")
        r = self.client.post("/api/corrections/", {"employee": self.emp.pk, "work_date": str(self.d), "kind": "set_status",
                             "new_status": "sick", "reason": "Certificat reçu"}, content_type=J, **hr)
        self.assertEqual(self.client.post(f"/api/corrections/{r.json()['id']}/reject/", {}, content_type=J, **mgr).status_code, 400)
        self.assertEqual(self.client.post(f"/api/corrections/{r.json()['id']}/reject/", {"note": "Non justifié"}, content_type=J, **mgr).status_code, 200)


class ReportTests(Base):
    def setUp(self):
        super().setUp()
        self.punch(self.emp, 8, 30)
        self.punch(self.emp, 17)
        self.punch(self.emp2, 8)
        recompute_day(self.emp, self.d)
        recompute_day(self.emp2, self.d)
        self.h = auth("hrm", "hr_manager")
        self.q = f"from={self.d}&to={self.d}"

    def test_daily_json_and_exports(self):
        j = self.client.get(f"/api/reports/daily/?{self.q}", **self.h).json()
        self.assertEqual(j["total"], 2)
        row = next(r for r in j["rows"] if r[1] == "E1")
        self.assertEqual(row[8], 30)  # late minutes
        csv = self.client.get(f"/api/reports/daily/?{self.q}&fmt=csv", **self.h)
        self.assertIn("text/csv", csv["Content-Type"])
        self.assertIn("Garnier", csv.content.decode("utf-8-sig"))
        x = self.client.get(f"/api/reports/monthly/?{self.q}&fmt=xlsx", **self.h)
        ws = load_workbook(io.BytesIO(x.content)).active
        self.assertEqual(ws["A4"].value, "Matricule")
        pdf = self.client.get(f"/api/reports/anomalies/?{self.q}&fmt=pdf", **self.h)
        self.assertEqual((pdf["Content-Type"], pdf.content[:4]), ("application/pdf", b"%PDF"))

    def test_filters_scope_and_access(self):
        self.assertEqual(self.client.get(f"/api/reports/late/?{self.q}&org_unit={self.other.pk}", **self.h).json()["total"], 0)
        head = auth("dh", "dept_head", [self.unit])
        j = self.client.get(f"/api/reports/daily/?{self.q}", **head).json()
        self.assertEqual([r[1] for r in j["rows"]], ["E1"])
        self.assertEqual(self.client.get(f"/api/reports/devices/?{self.q}", **head).status_code, 403)
        emp_only = auth("mob", "employee", employee=self.emp)
        self.assertEqual(self.client.get(f"/api/reports/daily/?{self.q}", **emp_only).status_code, 403)
        self.assertEqual(self.client.get("/api/reports/nope/", **self.h).status_code, 404)
        self.assertEqual(self.client.get("/api/reports/daily/?from=2026-03-05&to=2026-03-01", **self.h).status_code, 400)

    def test_csv_formula_injection_neutralised(self):
        self.emp.last_name = "=HYPERLINK(\"http://x\")"
        self.emp.save()
        csv = self.client.get(f"/api/reports/daily/?{self.q}&fmt=csv", **self.h).content.decode("utf-8-sig")
        self.assertNotIn(";=HYPERLINK", csv)

    def test_dashboard_numbers_come_from_data(self):
        j = self.client.get("/api/dashboard/", **self.h).json()
        self.assertEqual(j["headcount"], 2)
        head = auth("dh", "dept_head", [self.unit])
        self.assertEqual(self.client.get("/api/dashboard/", **head).json()["headcount"], 1)
        self.assertNotIn("devices", self.client.get("/api/dashboard/", **head).json())


class EmployeeIOTests(Base):
    def test_import_export_roundtrip_and_errors(self):
        dev = Device.objects.create(name="Entrée", vendor="zkteco", ip_address="10.0.0.2")
        PunchEvent.objects.create(device=dev, device_user_id="42", timestamp=timezone.now())
        csv = ("matricule;nom;prénom;unité;fonction;date_recrutement;statut;pointeuse;id_pointeuse\n"
               "M1;Durand;Anne;Direction RH > Paie;Comptable;15/01/2024;actif;Entrée;42\n"
               ";SansMatricule;X;;;;;;\n"
               "M3;Bad;Date;;;pas-une-date;actif;;\n").encode()
        h = auth("hrs", "hr_staff")
        r = self.client.post("/api/employees/import/", {"file": SimpleUploadedFile("p.csv", csv)}, **h).json()
        self.assertEqual((r["created"], len(r["errors"])), (1, 2))
        e = Employee.objects.get(matricule="M1")
        self.assertEqual((e.org_unit.name, e.org_unit.parent.name, str(e.hire_date)), ("Paie", "Direction RH", "2024-01-15"))
        self.assertEqual(DeviceUserMap.objects.get().device_user_id, "42")
        self.assertEqual(PunchEvent.objects.get().employee, e)  # existing raw punch linked retroactively
        out = self.client.get("/api/employees/export/", **h).content.decode("utf-8-sig")
        self.assertIn("M1;Durand;Anne;Direction RH > Paie", out)
        self.assertEqual(self.client.get("/api/employees/export/", **auth("v", "viewer", [self.unit])).status_code, 403)

    def test_dry_run_changes_nothing(self):
        csv = "matricule;nom\nZ9;Test\n".encode()
        r = self.client.post("/api/employees/import/?dry_run=1", {"file": SimpleUploadedFile("p.csv", csv)}, **auth("hrs", "hr_staff")).json()
        self.assertEqual(r["created"], 1)
        self.assertFalse(Employee.objects.filter(matricule="Z9").exists())


class MobileTests(Base):
    def setUp(self):
        super().setUp()
        self.h = auth("sophie", "employee", employee=self.emp)

    def test_disabled_by_default(self):
        self.assertEqual(self.client.post("/api/me/punch/", **self.h).status_code, 403)

    @override_settings(MOBILE_PUNCH_ENABLED=True)
    def test_punch_uses_server_time_and_rate_limits(self):
        r = self.client.post("/api/me/punch/", {"timestamp": "2000-01-01T00:00:00Z"}, content_type=J, **self.h)
        self.assertEqual(r.status_code, 201)
        p = PunchEvent.objects.get(source="mobile")
        self.assertLess(abs((timezone.now() - p.timestamp).total_seconds()), 5)
        self.assertEqual(self.client.post("/api/me/punch/", **self.h).status_code, 429)
        me = self.client.get("/api/me/", **self.h).json()
        self.assertEqual(me["employee"]["matricule"], "E1")

    @override_settings(MOBILE_PUNCH_ENABLED=True, MOBILE_PUNCH_ALLOWED_CIDRS=["10.1.0.0/16"])
    def test_network_restriction(self):
        self.assertEqual(self.client.post("/api/me/punch/", REMOTE_ADDR="8.8.8.8", **self.h).status_code, 403)
        self.assertEqual(self.client.post("/api/me/punch/", REMOTE_ADDR="10.1.2.3", **self.h).status_code, 201)

    def test_employee_role_sees_nothing_else_and_user_without_employee_rejected(self):
        self.assertEqual(self.client.get("/api/employees/", **self.h).status_code, 403)
        self.assertEqual(self.client.get("/api/punches/", **self.h).status_code, 403)
        self.assertEqual(self.client.get("/api/me/", **auth("nolink", "viewer")).status_code, 403)

    def test_history_and_planning(self):
        self.punch(self.emp, 8)
        self.punch(self.emp, 17)
        recompute_day(self.emp, self.d)
        j = self.client.get(f"/api/me/history/?from={self.d}&to={self.d}", **self.h).json()
        self.assertEqual((len(j["days"]), j["totals"]["worked_min"]), (1, 480))
        self.assertEqual(len(self.client.get("/api/me/planning/", **self.h).json()["days"]), 14)
